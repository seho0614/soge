"""v10.31 매트릭스 · 검정 엑셀 — 화면 · 페이지 · 백엔드 · API · DB(표 · 뷰 · 함수) · 연동 + 이번 판의 전수 검정 결과.

  python3 build_xlsx_v1031.py          → /home/claude/work/v3/out/SGB_Matrix_Verification_20261010_v10_31.xlsx   (XLSX_OUT=… 로 다른 곳에)

목록은 sgb/docs/matrix/matrix.json(`npm run matrix` 가 코드 · 실제 DB 에서 다시 만든 것)에서, 함수 · 연동은 개발 DB 에서 바로 읽습니다.
「이번 판」 칸 = v10.30 바탕(git 40ac141)의 matrix.json 에 없던 것(새로 생김).
"""
import json, os, re, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

K = '/home/claude/pkg/SGB_DevDocker_Integrated_20261002'
Q = os.path.dirname(os.path.abspath(__file__))
VER, VDATE = 'v10.31', '20261010'
OUT = os.environ.get('XLSX_OUT') or f"/home/claude/work/v3/out/SGB_Matrix_Verification_{VDATE}_{VER.replace('.', '_')}.xlsx"
DB = os.environ.get('DB', 'postgresql://sgb:sgb@localhost:5432/sgb')
BASE_COMMIT = os.environ.get('BASE_COMMIT', '40ac141')
m = json.load(open(K + '/sgb/docs/matrix/matrix.json', encoding='utf8'))
try: old = json.loads(subprocess.run(['git', '-C', K, 'show', f'{BASE_COMMIT}:sgb/docs/matrix/matrix.json'], capture_output=True, text=True).stdout)
except Exception: old = {'screens': [], 'pages': [], 'apis': [], 'db': [], 'summary': {}}
J = lambda f, d=None: (json.load(open(os.path.join(Q, f), encoding='utf8')) if os.path.exists(os.path.join(Q, f)) else d)
def rows(sql):
    p = subprocess.run(['psql', DB, '-tA', '-F', '\x1f', '-c', sql], capture_output=True, text=True)
    return [l.split('\x1f') for l in p.stdout.split('\n') if l.strip()]

wb = Workbook(); wb.remove(wb.active)
HDR = PatternFill('solid', fgColor='1F3A5F'); HF = Font(bold=True, color='FFFFFF')
OKF = PatternFill('solid', fgColor='E2F0D9'); NGF = PatternFill('solid', fgColor='FBE3E4'); WNF = PatternFill('solid', fgColor='FFF2CC'); NEWF = PatternFill('solid', fgColor='DDEBF7')
def cell(v):
    if isinstance(v, (list, tuple)): v = ', '.join(map(str, v))
    if isinstance(v, dict): v = json.dumps(v, ensure_ascii=False)
    if isinstance(v, str) and len(v) > 2000: v = v[:2000] + ' …'
    return v
def sheet(title, heads, rws, widths=None, status=None, mark=None):
    ws = wb.create_sheet(title[:31]); ws.append(heads)
    for c in ws[1]: c.fill = HDR; c.font = HF; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rws: ws.append([cell(v) for v in r])
    ws.freeze_panes = 'A2'
    if rws: ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(heads, 1): ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h) or min(60, max(10, len(str(h)) * 2 + 2))
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
        if status is not None:
            v = str(row[status].value or '')
            row[status].fill = OKF if v.startswith(('통과', '같음', '있음', '고침', '연결')) else NGF if v.startswith(('실패', '다름', '남음')) else WNF if v else PatternFill()
        if mark and mark(row):
            for c in row[:2]: c.fill = NEWF
    return ws

S = m['summary']; OS = old.get('summary', {})
NORM = lambda p: re.sub(r':[A-Za-z_]+|\{[^}]*\}', '{}', p or '')
old_api = {f"{a.get('method')} {NORM(a.get('path'))}" for a in old.get('apis', [])}
old_db = {d.get('name') for d in old.get('db', [])}
old_pg = {(p.get('app'), p.get('file')) for p in old.get('pages', [])}
fns = rows("""select p.proname, pg_get_function_identity_arguments(p.oid), pg_get_function_result(p.oid), regexp_replace(coalesce(obj_description(p.oid,'pg_proc'),''), '\\s+', ' ', 'g'),
                     (select count(*) from pg_trigger t where t.tgfoid = p.oid and not t.tgisinternal)
                from pg_proc p join pg_namespace n on n.oid = p.pronamespace where n.nspname = 'public' and p.prokind in ('f','p') order by 1, 2""")
mig = rows("select count(*) from schema_migrations") or [['']]

# ── 결과 읽기 ──
def read_final(f):
    out = []
    p = os.path.join(Q, f)
    if os.path.exists(p):
        for l in open(p, encoding='utf8').read().split('\n'):
            if ': ' in l and not l.startswith(('PART-', 'CHAIN')): k, v = l.split(': ', 1); out.append((k.strip(), v.strip()))
    return out
final = read_final('final_all.txt'); rer = {}
for k, v in read_final('final_all_reruns_1031.txt'): rer[k] = v          # 다시 돌린 줄이 앞섭니다(마지막 줄)
NOTE_RE = re.compile(r' — (?=시험을 고친|결함 고친|환경을 고친|새 시험|다시 빌드한|최신 소스|전 회차)')
def verdict(v):
    mm_ = NOTE_RE.search(v)
    if mm_: v = v[:mm_.start()]          # 다시 돌린 줄의 「— 무엇을 고친 뒤」 설명은 판정에서 뺌(설명 안의 「실패 13단계」 같은 글자)
    if v.startswith('건너뜀'): return '건너뜀'
    if re.search(r'남음|결정 필요', v): return '남음'
    if re.search(r'실패 [1-9]|FAIL|뚫림 [1-9]|exit=[1-9]|▲위험 [1-9]|Error|Traceback|Timeout|^실패|문제 [1-9]|다른 것 [1-9]', v): return '실패'
    mm = re.search(r'(\d+)\s*/\s*(\d+)', v)
    sk = re.search(r'건너뜀 (\d+)', v)   # 흐름 점검 · 시안 시험 — 「통과 a / b · 건너뜀 c」 는 a + c = b 면 통과(건너뜀은 따로 셈)
    if mm and sk and int(mm.group(1)) + int(sk.group(1)) == int(mm.group(2)): return '통과'
    if mm and mm.group(1) != mm.group(2) and not re.search(r'통과 \d+ / 실패 0|오류 0 / \d+', v): return '실패'
    if v in ('', 'done'): return '확인'
    return '통과'
last = {}; order = []
for k, v in final:                     # 같은 시험이 final_all.txt 에 두 번이면(구간을 다시 돌림) 나중 줄
    if k not in last: order.append(k)
    last[k] = v
seen = set(order); FIN = []
for k in order:
    v = last[k]; now = rer.get(k, v)
    FIN.append([k, verdict(now), now, v if k in rer else ''])
for k, v in rer.items():
    if k not in seen: FIN.append([k, verdict(v), v, '(전수 검정 밖에서 따로 돌림)'])
FIX = J('fixed_v1031.json', [])        # 고친 결함 · 고친 시험(검정 결과 문서와 같은 목록)
LEFT = J('left_v1031.json', [])        # 남은 것
V31 = J('srv_v1031.json', []); W5 = J('srv_w5.json', [])
LI = J('link_integrations.json', []); LC = J('link_check.json', []); DBL = J('db_link.json', [])

# ── 00 요약 ──
new_api = [a for a in m['apis'] if f"{a.get('method')} {NORM(a.get('path'))}" not in old_api]
new_db = [d for d in m['db'] if d.get('name') not in old_db]
new_pg = [p for p in m['pages'] if (p.get('app'), p.get('file')) not in old_pg]
cnt = lambda s: sum(1 for r in FIN if r[1] == s)
summary = [
 ['판', f'{VER} ({VDATE[:4]}-{VDATE[4:6]}-{VDATE[6:]})', '10/10 「시안 기준」 반영 다섯 묶음(w1 · w2 · w3 · w5 · w6) 합침 + 통합 처리'],
 ['화면(등록부)', S.get('screens_registry'), f"구현 연결 {S.get('screens_with_impl')} · 시안 연결 {S.get('screens_with_proto')} (v10.30: {OS.get('screens_registry')} · {OS.get('screens_with_impl')} · {OS.get('screens_with_proto')})"],
 ['페이지(사용자 웹 + 운영 화면)', f"{S.get('site_pages')} + {S.get('admin_pages')}", f"v10.30: {OS.get('site_pages')} + {OS.get('admin_pages')} · 새 페이지 {len(new_pg)}"],
 ['백엔드 처리기 = API', S.get('handlers'), f"v10.30: {OS.get('handlers')} · 새 API {len(new_api)} · OpenAPI {S.get('openapi_ops')}"],
 ['DB', f"표 {S.get('tables')} · 뷰 {S.get('views')} · 함수 {S.get('functions')}", f"v10.30: 표 {OS.get('tables')} · 뷰 {OS.get('views')} · 함수 {OS.get('functions')} · 마이그레이션 {S.get('migrations')}"],
 ['교차 점검(상 · 중 · 하)', f"{S.get('issues', {}).get('상')} · {S.get('issues', {}).get('중')} · {S.get('issues', {}).get('하')}", '0 · 0 · 0 이어야 함 — 처음 16 · 4 · 0 은 매트릭스 도구의 읽기 한계(07_미연결 · 11_고친것)'],
 ['외부 연동', f"{len(LI)}건 — " + ' · '.join(f"{x['code']} {x['status']}" for x in LI), '08_연동 시트'],
 ['전수 검정', f"{len(FIN)}줄 — 통과 {cnt('통과')} · 실패 {cnt('실패')} · 남음 {cnt('남음')} · 확인 {cnt('확인')} · 건너뜀 {cnt('건너뜀')}", '09_전수검정 시트(다시 돌린 시험은 다시 돌린 줄이 앞섬)'],
 ['새 API 역할별 시험(srv_v1031)', f"{sum(1 for r in V31 if r['ok'])} / {len(V31)}", '10_새API_역할별 시트'],
 ['고친 것 · 남은 것', f"고침 {len(FIX)} · 남음 {len(LEFT)}", '11 · 12 시트'],
]
_b = {k: J(f'btn_{k}.json') for k in ('site', 'bo')}
_BADB = ('죽은 단추', '누를 수 없음', '요청에 로그인 토큰이 없음', '다시 열었더니 없음')
for k, nm in (('site', '버튼 체크(사용자 웹 5역할)'), ('bo', '버튼 체크(운영 화면 역할별)')):
    if _b[k]: summary.append([nm, f"{'쪽' if k == 'site' else '상태'} {_b[k]['units']} · 단추 {sum(_b[k]['total'].values()):,} · 문제 {sum(_b[k]['total'].get(x, 0) for x in _BADB)}", '10d_확장점검 시트'])
_fl = []
for f in ('trade', 'mall', 'plaza', 'biz', 'proto'):
    d = J(f'flow_{f}.json', []); st = [x for fl in d for x in fl.get('steps', [])]; _fl.append(f"{f} {sum(1 for x in st if x.get('ok'))}/{len(st)}")
summary.append(['flow 점검', ' · '.join(_fl), '10e_flow단계 시트'])
sheet('00_요약', ['항목', '값', '비고'], summary, {'항목': 30, '값': 90, '비고': 80})

# ── 01 ~ 05 매트릭스 ──
old_sc = {s.get('id'): s for s in old.get('screens', [])}
def sc_note(s):
    o = old_sc.get(s.get('id'))
    if not o: return '새 화면'
    if (o.get('impl') or []) != (s.get('impl') or []): return '구현 파일 바뀜'
    if sorted(o.get('apis_called') or []) != sorted(s.get('apis_called') or []): return '부르는 API 바뀜'
    return ''
sheet('01_화면목록', ['화면 ID', '이름', '묶음', '주소', '역할', '게이트', '구현 파일', '연결 방법', '시안 화면', 'API(등록부)', 'API(코드)', '백엔드에 없는 API', '이번 판'],
      [[s.get('id'), s.get('name'), s.get('group'), s.get('url'), s.get('roles'), s.get('gates'), s.get('impl'), s.get('match'), s.get('proto'), s.get('apis_registry'), s.get('apis_called'), s.get('apis_missing_backend'), sc_note(s)] for s in m['screens']],
      {'화면 ID': 18, '이름': 26, '주소': 24, '구현 파일': 44, '시안 화면': 14, 'API(등록부)': 46, 'API(코드)': 46, '이번 판': 16}, mark=lambda r: bool(r[12].value))
sheet('02_페이지목록', ['앱', '파일', '주소 · 탭', '화면 ID', '부르는 API', '링크', '이번 판'],
      [[p.get('app'), p.get('file'), p.get('route'), p.get('ids'), p.get('calls'), p.get('links'), '' if (p.get('app'), p.get('file')) in old_pg else '새 페이지'] for p in m['pages']],
      {'파일': 52, '주소 · 탭': 28, '화면 ID': 30, '부르는 API': 80, '링크': 50}, mark=lambda r: bool(r[6].value))
sheet('03_백엔드목록', ['모듈', '메서드', '경로', '컨트롤러', '처리기', '파일', '역할', '게이트', '서비스 호출', '표'],
      [[b.get('module'), b.get('method'), b.get('path'), b.get('controller'), b.get('handler'), b.get('file'), b.get('roles'), b.get('gate'), b.get('calls') or b.get('services'), b.get('tables')] for b in m['backend']],
      {'경로': 40, '컨트롤러': 28, '처리기': 22, '파일': 44, '역할': 24, '서비스 호출': 40, '표': 40})
# 03b 서비스 — src/modules/*/*.service.ts(클래스 · 공개 메서드 · 줄 수) · 「이번 판」 = v10.30 바탕(git BASE_COMMIT)에 없던 파일
import glob as _glob
_base_files = set(subprocess.run(['git', '-C', K, 'ls-tree', '-r', '--name-only', BASE_COMMIT, 'sgb/src/modules'], capture_output=True, text=True).stdout.split())
_svc = []
for f in sorted(_glob.glob(K + '/sgb/src/modules/**/*.service.ts', recursive=True)):
    src = open(f, encoding='utf8').read(); rel = f[len(K) + 1:]
    cls = re.findall(r'export class (\w+)', src)
    meth = re.findall(r'^  (?:async )?([a-zA-Z_]\w*)\s*\([^)]*\)[^;{]*\{', src, re.M)
    meth = [x for x in meth if x not in ('constructor', 'if', 'for', 'while', 'switch', 'catch')]
    flag = '새 서비스' if rel not in _base_files else ('바뀜' if subprocess.run(['git', '-C', K, 'diff', '--quiet', BASE_COMMIT, '--', rel]).returncode != 0 else '')
    _svc.append([rel.split('/')[3], rel.split('/')[-1], ', '.join(cls), len(meth), ', '.join(meth[:40]) + (' …' if len(meth) > 40 else ''), src.count('\n') + 1, flag])
sheet('03b_백엔드_서비스', ['모듈', '파일', '클래스', '메서드 수', '메서드', '줄 수', '이번 판'], _svc, {'파일': 36, '클래스': 30, '메서드': 90, '이번 판': 12}, mark=lambda r: bool(r[6].value))
sheet('04_API목록', ['메서드', '경로', '모듈', '컨트롤러', '역할', '게이트', '상태', '화면(등록부)', '부르는 곳', '표', '이번 판'],
      [[a.get('method'), a.get('path'), a.get('module'), a.get('controller'), a.get('roles'), a.get('gate'), a.get('status'), a.get('screens_registry'), a.get('callers'), a.get('tables'),
        '' if f"{a.get('method')} {NORM(a.get('path'))}" in old_api else '새 API'] for a in m['apis']],
      {'경로': 40, '컨트롤러': 32, '역할': 22, '화면(등록부)': 30, '부르는 곳': 50, '표': 40}, mark=lambda r: bool(r[10].value))
sheet('05_DB_표뷰', ['스키마', '이름', '종류', '칸 수', '행 수(개발 DB)', '트리거', '쓰는 API 수', '모듈', '설명', '이번 판'],
      [[d.get('schema'), d.get('name'), d.get('kind'), d.get('cols'), d.get('rows'), d.get('triggers'), d.get('apis'), d.get('modules'), d.get('comment'), '' if d.get('name') in old_db else '새 표 · 뷰'] for d in m['db']],
      {'이름': 36, '모듈': 30, '설명': 70}, mark=lambda r: bool(r[9].value))
sheet('06_DB_함수', ['함수', '인자', '돌려주는 것', '트리거로 쓰임(수)', '설명'], [[f[0], f[1], f[2], f[4], f[3]] for f in fns], {'함수': 34, '인자': 50, '돌려주는 것': 24, '설명': 80})
# 06b 마이그레이션 — prisma/migrations 폴더 순서 · 개발 DB 의 적용 기록(schema_migrations) · 첫 설명 줄
_app = {r[0]: r[1:] for r in rows("select name, to_char(applied_at at time zone 'Asia/Seoul','YYYY-MM-DD HH24:MI'), checksum, coalesce(duration_ms::text,'') from schema_migrations")}
_base_mig = {x.split('/')[3] for x in subprocess.run(['git', '-C', K, 'ls-tree', '-r', '--name-only', BASE_COMMIT, 'sgb/prisma/migrations'], capture_output=True, text=True).stdout.split() if x.count('/') >= 4}
_mig = []
for d in sorted(os.listdir(K + '/sgb/prisma/migrations')):
    f = K + '/sgb/prisma/migrations/' + d + '/migration.sql'
    if not os.path.isfile(f): continue
    first = next((l.strip('- ').strip() for l in open(f, encoding='utf8') if l.strip().startswith('--') and re.search(r'[0-9A-Za-z가-힣]', l)), '')
    a = _app.get(d)
    _mig.append([d, a[0] if a else '적용 안 됨', a[1] if a else '', a[2] if a else '', os.path.getsize(f), first[:200], '' if d in _base_mig else '새 마이그레이션'])
sheet('06b_DB_마이그레이션', ['이름', '개발 DB 적용 시각(한국)', '체크섬', '걸린 ms', '크기(바이트)', '첫 설명 줄', '이번 판'], _mig, {'이름': 40, '첫 설명 줄': 90, '이번 판': 16}, mark=lambda r: bool(r[6].value))
sheet('07_미연결', ['등급', '영역', '대상', '내용'], [[i.get('sev'), i.get('area'), i.get('what'), i.get('detail')] for i in (m.get('issues') or [])] or [['—', '—', '없음', '화면 · API · 백엔드 · DB 사이에 끊긴 곳 0 (상 0 · 중 0 · 하 0)']], {'대상': 40, '내용': 100})

# ── 08 연동 ──
ENVS = {'TOSS': 'TOSS_SECRET_KEY · TOSS_API_BASE (가짜 결제사 3129 / srv_bill_toss 3139)', 'S3': 'S3_BUCKET · S3_ENDPOINT · AWS_* (가짜 S3 9022)', 'PAYOUT_VERIFY': '1원 인증(가짜 13023)', 'BROKER': '국토부 중개사 대조(가짜 13023 /broker)', 'NTS': '국세청 상태 조회(가짜 13023 /nts)'}
integ = rows("select code, name, status, coalesce(secret_key_name,''), secret_registered::text, blocks_open::text, regexp_replace(coalesce(memo,''), '\\s+', ' ', 'g') from integrations order by sort_order, code")
lk = {x['code']: x for x in LI}
sheet('08_연동', ['코드', '이름', '상태', '환경값(키 이름)', '서버 코드가 그 키를 읽음', '키 등록', '오픈 차단', '시험 방법(가짜 서버)', '비고'],
      [[r[0], r[1], r[2], r[3], {True: '읽음', False: '안 읽음', None: '—'}.get((lk.get(r[0]) or {}).get('code_reads_key'), '—'), r[4], r[5], ENVS.get(r[0], '—'), r[6]] for r in integ],
      {'이름': 26, '환경값(키 이름)': 28, '시험 방법(가짜 서버)': 46, '비고': 50})

# ── 09 전수 검정 ──
sheet('09_전수검정', ['시험', '판정', '결과 줄', '처음 결과(다시 돌린 시험만)'], FIN, {'시험': 22, '결과 줄': 130, '처음 결과(다시 돌린 시험만)': 80}, status=1)
# ── 10 새 API 역할별 · W5 ──
sheet('10_새API_역할별', ['단계', '판정', '자세히'], [[r['step'], '건너뜀' if r.get('skip') else '통과' if r['ok'] else '실패', r.get('detail', '')] for r in V31], {'단계': 110, '자세히': 90}, status=1)
sheet('10b_파트너센터_W5', ['단계', '판정', '자세히'], [[r['step'], '통과' if r['ok'] else '실패', r.get('detail', '') if not r['ok'] else ''] for r in W5], {'단계': 110, '자세히': 60}, status=1)
sheet('10c_연동·DB연동', ['시험', '단계', '판정', '자세히'], [['link_check', r['step'], '건너뜀' if r.get('skip') else '통과' if r['ok'] else '실패', r.get('detail', '') if not r['ok'] else ''] for r in LC] +
      [['db_link', r['step'], '건너뜀' if r.get('skip') else '통과' if r['ok'] else '실패', r.get('detail', '') if not r['ok'] else ''] for r in DBL], {'단계': 110, '자세히': 60}, status=2)
# ── 10d · 10e 확장 점검(버튼 · flow) ──
BAD_BTN = ('죽은 단추', '누를 수 없음', '요청에 로그인 토큰이 없음', '다시 열었더니 없음')
ext = []
for kind in ('site', 'bo'):
    b = J(f'btn_{kind}.json')
    if b:
        bad = sum(b['total'].get(k, 0) for k in BAD_BTN)
        ext.append([f'버튼 체크 — {"사용자 웹(5역할)" if kind == "site" else "운영 화면(역할별)"}', '통과' if bad == 0 else '실패',
                    f"{'쪽' if kind == 'site' else '상태'} {b['units']} · 단추 {sum(b['total'].values())} · " + ' · '.join(f'{k} {v}' for k, v in sorted(b['total'].items(), key=lambda x: -x[1])) + f' · 문제 {bad}'])
FLOWS = []
for f in ('trade', 'mall', 'plaza', 'biz', 'proto'):
    d = J(f'flow_{f}.json', [])
    st = [x for fl in d for x in fl.get('steps', [])]
    okn = sum(1 for x in st if x.get('ok') and not x.get('skip')); skn = sum(1 for x in st if x.get('skip'))
    ext.append([f'flow_{f}', '통과' if st and okn + skn == len(st) else '실패', f'흐름 {len(d)} · 단계 통과 {okn} / {len(st)}' + (f' · 건너뜀 {skn}' if skn else '')])
    FLOWS += [[f'flow_{f}', fl.get('flow'), x.get('no'), x.get('step'), x.get('how'), '건너뜀' if x.get('skip') else '통과' if x.get('ok') else '실패', x.get('detail')] for fl in d for x in fl.get('steps', [])]
for k in ('crawl_site', 'crawl_admin', 'link_check', 'db_link'):
    v = rer.get(k) or last.get(k)
    if v and NOTE_RE.search(v): v = v[:NOTE_RE.search(v).start()]
    if v: ext.append([k, verdict(v) if k not in ('crawl_site', 'crawl_admin') else ('통과' if re.search(r'실패 0|스크립트 오류 0쪽 · API 실패 0쪽 · 오류 문구 0쪽', v) else '실패'), v])
sheet('10d_확장점검', ['점검', '판정', '결과'], ext or [['—', '', '']], {'점검': 34, '결과': 140}, status=1)
sheet('10e_flow단계', ['흐름 묶음', '흐름', '번호', '단계', '확인 방법', '판정', '자세히'], FLOWS or [['—'] * 7], {'흐름': 36, '단계': 90, '자세히': 70}, status=5)
# ── 11 · 12 ──
_TF = []
for l in open(os.path.join(Q, 'final_all_reruns_1031.txt'), encoding='utf8').read().split('\n'):
    if ': ' not in l: continue
    k_, v_ = l.split(': ', 1); mm_ = re.search(r' — (?=시험을 고친|결함 고친|환경을 고친|새 시험|다시 빌드한|최신 소스|전 회차)', v_)
    if not mm_: continue
    for sg in v_[mm_.end():].split(' · 시험을 고친')[1:] if not v_[mm_.end():].startswith('시험을 고친') else [v_[mm_.end():][len('시험을 고친'):]]:
        _TF.append([f'T-{len(_TF) + 1}', '시험', k_.strip(), '시험이 v10.31 의 새 화면 · 정책을 못 따라감', '시험을 고친' + sg, 'qa/' + k_.strip() + '.py', v_[:mm_.start()][:80]])
sheet('11_고친것', ['번호', '구분', '무엇', '증상', '한 일', '파일 · 커밋', '다시 돌린 결과'], ([[f.get('no'), f.get('kind'), f.get('what'), f.get('symptom'), f.get('done'), f.get('where'), f.get('rerun')] for f in FIX] + _TF) or [['—'] * 7],
      {'무엇': 40, '증상': 60, '한 일': 60, '파일 · 커밋': 40, '다시 돌린 결과': 30})
sheet('12_남은것', ['번호', '구분', '무엇', '까닭', '필요한 결정 · 할 일'], [[f.get('no'), f.get('kind'), f.get('what'), f.get('why'), f.get('todo')] for f in LEFT] or [['—', '', '없음', '', '']],
      {'무엇': 50, '까닭': 70, '필요한 결정 · 할 일': 60})

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('저장', OUT, '· 시트', len(wb.sheetnames), wb.sheetnames)
