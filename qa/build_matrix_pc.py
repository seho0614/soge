"""집 PC 매트릭스 엑셀 — 화면 · 페이지 · 백엔드 · API · DB · 함수 · 연동 + realcheck 결과  (v10.43 · 2026-10-15)

  qa\\matrix_pc.ps1 이 부릅니다(파이썬 컨테이너 안). 직접 돌릴 때:
    python build_matrix_pc.py matrix_out

  입력 (matrix_out/)
    matrix.json        npm run matrix(scripts/gen-matrix.js) 가 코드 · 개발 DB 에서 만든 것 — 화면 · 페이지 · 백엔드 · API · DB · 점검
    functions.csv      개발 DB 의 함수 목록(psql)
    integrations.csv   개발 DB 의 외부 연동 표(psql)
    VERSION            kit 판
    realcheck_*.txt    (있으면) 가장 최근 realcheck 결과
  출력
    matrix_out/SGB_Matrix_Verification_<날짜>_<판>.xlsx

  kit 작업 환경용 build_xlsx_v10xx.py 와 달리 그 환경의 검정 결과 파일(btn_* · flow_* · final_all_reruns_*)을 요구하지 않습니다.
"""
import csv, glob, json, os, sys
from datetime import datetime, timedelta, timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

D = sys.argv[1] if len(sys.argv) > 1 else 'matrix_out'
m = json.load(open(os.path.join(D, 'matrix.json'), encoding='utf-8'))
VER = open(os.path.join(D, 'VERSION'), encoding='utf-8').read().strip() if os.path.exists(os.path.join(D, 'VERSION')) else 'v?'
KST = timezone(timedelta(hours=9))
DAY = datetime.now(KST).strftime('%Y%m%d')
OUT = os.path.join(D, f"SGB_Matrix_Verification_{DAY}_{VER.replace('.', '_')}.xlsx")


def rcsv(name):
    p = os.path.join(D, name)
    if not os.path.exists(p):
        return [], []
    rows = list(csv.reader(open(p, encoding='utf-8-sig')))
    return (rows[0], rows[1:]) if rows else ([], [])


HDR = PatternFill('solid', fgColor='1F3A5F'); HF = Font(bold=True, color='FFFFFF')
OKF = PatternFill('solid', fgColor='E2F0D9'); NGF = PatternFill('solid', fgColor='FBE3E4'); WNF = PatternFill('solid', fgColor='FFF2CC')
wb = Workbook(); wb.remove(wb.active)


def cell(v):
    if isinstance(v, (list, tuple)): v = ' | '.join(map(str, v))
    if isinstance(v, dict): v = json.dumps(v, ensure_ascii=False)
    if isinstance(v, str) and len(v) > 2000: v = v[:2000] + ' …'
    return v


def sheet(title, heads, rows, status=None, good=(), bad=()):
    ws = wb.create_sheet(title[:31]); ws.append(heads)
    for c in ws[1]: c.fill = HDR; c.font = HF; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows: ws.append([cell(v) for v in r])
    ws.freeze_panes = 'A2'
    if rows: ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(heads, 1): ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(str(h)) * 2 + 2))
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
        if status is not None:
            v = str(row[status].value or '')
            row[status].fill = OKF if v.startswith(good) else NGF if v.startswith(bad) else WNF if v else PatternFill()
    return ws


def table(title, key, cols, status=None, good=(), bad=()):
    rows = [[r.get(c, '') for c in cols] for r in m.get(key, [])]
    sheet(title, cols, rows, status, good, bad)
    return len(rows)


# realcheck
rc_files = sorted(glob.glob(os.path.join(D, 'realcheck_*.txt')))
rc_lines = open(rc_files[-1], encoding='utf-8').read().splitlines() if rc_files else []
rc_rows = [[l.strip()[:2], l.strip()[2:].strip()] for l in rc_lines if l.strip().startswith(('통과', '실패'))]
rc_rows += [['요약', l.strip()] for l in rc_lines if l.strip().startswith(('통과 ', '안전 ', 'SQL 대조', '진짜 이미지 확인'))]

S = m.get('summary', {})
fh, fr = rcsv('functions.csv')
ih, ir = rcsv('integrations.csv')
summ = [
    ['판', VER], ['만든 때(KST)', datetime.now(KST).strftime('%Y-%m-%d %H:%M')], ['matrix.json 만든 때', S.get('generated_at', '')],
    ['마이그레이션', S.get('migrations', '')],
    ['화면(등록) · 구현 있음 · 시안 있음', f"{S.get('screens_registry', '')} · {S.get('screens_with_impl', '')} · {S.get('screens_with_proto', '')}"],
    ['페이지 사용자 웹 · 운영 화면', f"{S.get('site_pages', '')} · {S.get('admin_pages', '')}"],
    ['백엔드 핸들러 · OpenAPI', f"{S.get('handlers', '')} · {S.get('openapi_ops', '')}"],
    ['DB 표 · 뷰 · 함수', f"{S.get('tables', '')} · {S.get('views', '')} · {S.get('functions', '') or len(fr)}"],
    ['외부 연동', len(ir)],
    ['교차 점검 상 · 중 · 하', ' · '.join(str(S.get('issues', {}).get(k, '')) for k in ('상', '중', '하'))],
    ['realcheck', (rc_files and os.path.basename(rc_files[-1])) or '없음'],
    ['realcheck 결과', next((l.strip() for l in rc_lines if l.strip().startswith('진짜 이미지 확인')), '없음')],
]
sheet('요약', ['항목', '값'], summ)
table('화면', 'screens', ['id', 'name', 'group', 'url', 'bo', 'roles', 'gates', 'impl', 'match', 'proto', 'apis_registry', 'apis_called', 'apis_missing_backend'])
table('페이지', 'pages', ['app', 'route', 'file', 'ids', 'calls', 'links'])
table('백엔드', 'backend', ['module', 'controller', 'handler', 'method', 'path', 'roles', 'gate', 'services', 'tables', 'fns', 'file'])
table('API', 'apis', ['method', 'path', 'module', 'controller', 'roles', 'gate', 'status', 'screens_registry', 'callers', 'tables', 'table_scope', 'fns'])
table('DB', 'db', ['schema', 'name', 'kind', 'cols', 'rows', 'triggers', 'apis', 'modules', 'comment'])
sheet('DB 함수', fh or ['함수'], fr)
sheet('연동', ih or ['code'], ir, status=(ih.index('status') if 'status' in ih else None), good=('LIVE',), bad=('NOT_STARTED', 'FAILED'))
table('교차 점검', 'issues', ['sev', 'area', 'what', 'detail'], status=0, good=('하',), bad=('상',))
sheet('realcheck', ['결과', '내용'], rc_rows, status=0, good=('통과', '요약'), bad=('실패',))

wb.save(OUT)
print(f'만듦 {OUT}  (화면 {len(m.get("screens", []))} · API {len(m.get("apis", []))} · DB {len(m.get("db", []))} · 함수 {len(fr)} · 연동 {len(ir)} · realcheck {len(rc_rows)}줄)')
